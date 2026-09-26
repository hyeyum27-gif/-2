"""PQ4R 학습노트(편집 가능한 Word DOCX, A4 세로 5쪽) 생성기.

사용법:
    python3 build_pq4r.py units/elastic_force.py

소단원 내용은 units/*.py 파일에만 적는다. 이 파일의 디자인(색상, 표 너비,
행 높이, fixed layout, tblGrid 처리)은 특별한 요청이 없으면 수정하지 않는다.
"""
import importlib.util
import os
import sys

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = os.path.dirname(os.path.abspath(__file__))
FONT = "Noto Sans KR"  # Windows 호환 우선 시 "맑은 고딕" (줄바꿈·표 높이 재검수 필요)

# 단계별 색상
P_C, P_PALE = "138BC6", "DCEFF7"      # Preview 파랑
Q_C, Q_PALE = "FF9F28", "FFF1D9"      # Question 주황
R1_C, R1_PALE = "65B33F", "EDF7DF"    # Read 초록
R2_C, R2_PALE = "FF8A68", "FFE2D8"    # Reflect 코랄
R3_C, R3_PALE = "72269E", "F3E6FA"    # Recite 보라
R4_C, R4_PALE = "2BA89E", "DDF6F3"    # Review 청록
TITLE_C = "55C7E7"                    # PQ4R 제목 하늘색
SCI_FILL, SCI_LINE = "FFE49A", "F4B848"  # 과학 표시
TEXT_C, GRAY_C, LINE_C = "333333", "777777", "BFBFBF"

PAGE_W = 186  # 본문 너비(mm) = 210 - 좌우 여백 12*2

doc = Document()


# ---------------------------------------------------------------- 기본 함수
def shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def border(cell, color=LINE_C, size=6, sides=("top", "left", "bottom", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcBorders"))
    if old is not None:
        tcPr.remove(old)
    b = OxmlElement("w:tcBorders")
    tcPr.append(b)
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}")
        if s in sides:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(size))
            e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)


def margins(cell, top=1.5, left=2.5, bottom=1.5, right=2.5):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(int(v * 56.7)))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def _font(run, size, bold, color):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def write(cell, text, size=10.5, bold=False, color=TEXT_C, align="left",
          valign="center", spacing=1.25):
    """셀 내용을 text로 교체한다. text의 줄바꿈(\\n)은 문단으로 나눈다."""
    cell.vertical_alignment = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                               "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER}[valign]
    lines = text.split("\n")
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = spacing
        if line:
            _font(p.add_run(line), size, bold, color)
    return p


def width(cell, mm):
    cell.width = Mm(mm)
    tcPr = cell._tc.get_or_add_tcPr()
    w = tcPr.find(qn("w:tcW"))
    if w is None:
        w = OxmlElement("w:tcW")
        tcPr.insert(0, w)
    w.set(qn("w:w"), str(int(mm * 56.7)))
    w.set(qn("w:type"), "dxa")


def table(rows, cols, widths):
    """열 너비가 Word·LibreOffice 모두에서 유지되도록
    columns.width + tblGrid/gridCol + 셀 tcW + tblLayout=fixed 를 모두 지정한다."""
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths) * 56.7)))
    tw.set(qn("w:type"), "dxa")
    old = tblPr.find(qn("w:tblW"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(tw)
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 56.7)))
    for i, w in enumerate(widths):
        t.columns[i].width = Mm(w)
        for c in t.columns[i].cells:
            width(c, w)
            margins(c)
            border(c)
    return t


def height(row, mm, rule=WD_ROW_HEIGHT_RULE.EXACTLY):
    row.height = Mm(mm)
    row.height_rule = rule
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)


def spacer(mm):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Mm(mm)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.add_run().font.size = Pt(1)


def page_break():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    r = p.add_run()
    r.font.size = Pt(1)
    r.add_break(WD_BREAK.PAGE)


def science_badge(cell, size=13):
    shading(cell, SCI_FILL)
    border(cell, SCI_LINE, 12)
    write(cell, "과학", size, True, "7A4B00", "center")


def titlebar(subtitle, pages):
    t = table(1, 3, [120, 36, 30])
    height(t.rows[0], 17)
    a, b, c = t.row_cells(0)
    for x in (a, b):
        border(x, sides=())
    write(a, "PQ4R 노트", 26, True, TITLE_C)
    write(b, "", 10)
    science_badge(c)
    spacer(2)
    t = table(1, 2, [142, 44])
    height(t.rows[0], 12)
    a, b = t.row_cells(0)
    border(a, TITLE_C, 12, sides=("bottom",))
    border(b, TITLE_C, 12, sides=("bottom",))
    write(a, subtitle, 15, True)
    write(b, f"교과서 {pages}", 11, True, GRAY_C, "right")


def small_header(subtitle, pages):
    t = table(1, 3, [124, 38, 24])
    height(t.rows[0], 11)
    a, b, c = t.row_cells(0)
    border(a, TITLE_C, 12, sides=("bottom",))
    border(b, TITLE_C, 12, sides=("bottom",))
    write(a, subtitle, 13, True)
    write(b, f"교과서 {pages}", 10, True, GRAY_C, "right")
    science_badge(c, 11)


def band(letter, english, phrase, color, pale):
    t = table(1, 3, [18, 42, 126])
    height(t.rows[0], 13)
    a, b, c = t.row_cells(0)
    for x in (a, b):
        shading(x, color)
        border(x, color, 6)
    shading(c, pale)
    border(c, color, 6)
    write(a, letter, 20, True, "FFFFFF", "center")
    write(b, english, 15, True, "FFFFFF")
    write(c, phrase, 10.5, True, color)


def head_row(t, labels, color, pale):
    height(t.rows[0], 9)
    for cell, text in zip(t.rows[0].cells, labels):
        shading(cell, pale)
        border(cell, color, 8)
        write(cell, text, 10.5, True, color, "center")


def body_cell(cell, color):
    border(cell, color, 6)


# ---------------------------------------------------------------- 내용 불러오기
if len(sys.argv) != 2:
    sys.exit("사용법: python3 build_pq4r.py units/<소단원>.py")
spec = importlib.util.spec_from_file_location("unit", sys.argv[1])
U = importlib.util.module_from_spec(spec)
spec.loader.exec_module(U)
for name, n in (("PREVIEW", 4), ("QUESTIONS", 3), ("KEYWORDS", 4),
                ("REFLECT", 4), ("RECITE", 3), ("REVIEW", 3)):
    assert len(getattr(U, name)) == n, f"{name} 항목은 {n}개여야 합니다."
OUT = os.path.join(ROOT, "output", U.OUT)

sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(12)
sec.left_margin = sec.right_margin = Mm(12)
sec.header_distance = sec.footer_distance = Mm(6)
st = doc.styles["Normal"]
st.font.name = FONT
st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(0)

# ---------------------------------------------------------------- 1쪽 Preview
titlebar(U.TITLE, U.PAGES)
spacer(2)
t = table(1, 1, [PAGE_W])
height(t.rows[0], 15)
shading(t.cell(0, 0), "F7F7F7")
border(t.cell(0, 0), LINE_C, 4)
write(t.cell(0, 0),
      "교과서를 직접 읽으며 질문의 답을 찾아 쓰는 자기주도 학습노트입니다.\n"
      "P → Q → R(Read) → R(Reflect) → R(Recite) → R(Review) 순서로 차근차근 채워 보세요.",
      9.5, False, "555555")
spacer(3)
band("P", "Preview", "교과서 내용과 학습 목표를 먼저 훑어보아요.", P_C, P_PALE)
spacer(2)
t = table(2, 1, [PAGE_W])
height(t.rows[0], 9)
height(t.rows[1], 17)
shading(t.cell(0, 0), P_PALE)
border(t.cell(0, 0), P_C, 8)
border(t.cell(1, 0), P_C, 6)
write(t.cell(0, 0), f"교과서 확인  |  {U.PAGES}", 10.5, True, P_C)
write(t.cell(1, 0), U.PREVIEW_CHECK, 10, False, TEXT_C, spacing=1.4)
spacer(3)
t = table(5, 2, [48, 138])
head_row(t, ["항목", "학습 목표 및 핵심 내용"], P_C, P_PALE)
for i, item in enumerate(U.PREVIEW, 1):
    height(t.rows[i], 40)
    a, b = t.row_cells(i)
    shading(a, P_PALE)
    border(a, P_C, 6)
    body_cell(b, P_C)
    write(a, item, 11, True, P_C, "center")
page_break()

# ---------------------------------------------------------------- 2쪽 Question
small_header(U.TITLE, U.PAGES)
spacer(3)
band("Q", "Question", "교과서에서 답을 찾을 핵심 질문을 만들어요.", Q_C, Q_PALE)
spacer(3)
t = table(4, 3, [14, 70, 102])
head_row(t, ["번호", "질문", "교과서에서 찾은 답"], Q_C, Q_PALE)
for i, q in enumerate(U.QUESTIONS, 1):
    height(t.rows[i], 76)
    a, b, c = t.row_cells(i)
    shading(a, Q_PALE)
    for x in (a, b, c):
        body_cell(x, Q_C)
    write(a, str(i), 16, True, Q_C, "center")
    write(b, q, 11, False, TEXT_C, valign="top", spacing=1.45)
page_break()

# ---------------------------------------------------------------- 3쪽 Read
small_header(U.TITLE, U.PAGES)
spacer(3)
band("R", "Read", "교과서를 읽으며 핵심어의 뜻과 설명을 정리해요.", R1_C, R1_PALE)
spacer(3)
t = table(5, 2, [46, 140])
head_row(t, ["핵심어", "교과서에서 찾은 뜻·설명"], R1_C, R1_PALE)
for i, k in enumerate(U.KEYWORDS, 1):
    height(t.rows[i], 57)
    a, b = t.row_cells(i)
    shading(a, R1_PALE)
    body_cell(a, R1_C)
    body_cell(b, R1_C)
    write(a, k, 12, True, R1_C, "center")
page_break()

# ---------------------------------------------------------------- 4쪽 Reflect
small_header(U.TITLE, U.PAGES)
spacer(3)
band("R", "Reflect", "탐구·그림·표·그래프·생활 사례와 핵심 개념을 연결해요.", R2_C, R2_PALE)
spacer(3)
t = table(5, 2, [74, 112])
head_row(t, ["연결 활동", "나의 생각"], R2_C, R2_PALE)
for i, r in enumerate(U.REFLECT, 1):
    height(t.rows[i], 57)
    a, b = t.row_cells(i)
    shading(a, R2_PALE)
    body_cell(a, R2_C)
    body_cell(b, R2_C)
    write(a, r, 10.5, False, TEXT_C, valign="top", spacing=1.45)
page_break()

# ---------------------------------------------------------------- 5쪽 Recite · Review
small_header(U.TITLE, U.PAGES)
spacer(3)
band("R", "Recite", "책을 덮고 자신의 말로 설명해 보아요.", R3_C, R3_PALE)
spacer(3)
t = table(4, 2, [74, 112])
head_row(t, ["회상 질문", "내 답"], R3_C, R3_PALE)
for i, r in enumerate(U.RECITE, 1):
    height(t.rows[i], 46)
    a, b = t.row_cells(i)
    shading(a, R3_PALE)
    body_cell(a, R3_C)
    body_cell(b, R3_C)
    write(a, r, 10.5, False, TEXT_C, valign="top", spacing=1.45)
spacer(7)
band("R", "Review", "오늘 배운 내용을 스스로 점검해요.", R4_C, R4_PALE)
spacer(3)
t = table(4, 3, [140, 23, 23])
head_row(t, ["자기 점검", "잘함", "보충"], R4_C, R4_PALE)
for i, r in enumerate(U.REVIEW, 1):
    height(t.rows[i], 18)
    a, b, c = t.row_cells(i)
    for x in (a, b, c):
        body_cell(x, R4_C)
    write(a, r, 10.5)
    write(b, "□", 13, False, R4_C, "center")
    write(c, "□", 13, False, R4_C, "center")

# ---------------------------------------------------------------- 메타데이터·저장
doc.core_properties.title = f"PQ4R 노트 - {U.TITLE}"
doc.core_properties.subject = f"미래엔 중등 과학 {U.TITLE} ({U.PAGES})"
doc.core_properties.author = "PQ4R 학습노트"
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT)
